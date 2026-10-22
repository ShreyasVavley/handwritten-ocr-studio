"""
Exporters for OCR text to Word (.docx), PDF (.pdf), CSV (.csv), and TXT (.txt).
"""

import io
import csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors

def export_to_txt(text_content):
    bio = io.BytesIO()
    bio.write(text_content.encode("utf-8"))
    bio.seek(0)
    return bio

def export_to_csv(text_content):
    bio = io.StringIO()
    writer = csv.writer(bio)
    writer.writerow(["Line Number", "Transcribed Text"])
    for idx, line in enumerate(text_content.splitlines(), start=1):
        if line.strip():
            writer.writerow([idx, line.strip()])
    
    bytes_io = io.BytesIO()
    bytes_io.write(bio.getvalue().encode("utf-8-sig"))
    bytes_io.seek(0)
    return bytes_io

def export_to_docx(text_content):
    doc = Document()
    
    # Title
    title = doc.add_heading("ScribeCraft • Transcribed Document", level=1)
    title.style.font.name = "Calibri"
    
    # Body text
    for line in text_content.splitlines():
        if line.strip():
            p = doc.add_paragraph(line.strip())
            p.style.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(6)

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

def export_to_pdf(text_content):
    bio = io.BytesIO()
    c = canvas.Canvas(bio, pagesize=letter)
    width, height = letter

    c.setFont("Helvetica-Bold", 16)
    c.drawString(54, height - 54, "ScribeCraft • Transcribed Document")
    c.setStrokeColor(colors.HexColor("#ede8df"))
    c.line(54, height - 64, width - 54, height - 64)

    c.setFont("Helvetica", 10)
    y = height - 90
    margin_bottom = 54
    line_height = 16

    for line in text_content.splitlines():
        if not line.strip():
            y -= line_height / 2
            continue
        
        # Word wrap if line is long
        words = line.split(" ")
        curr_line = ""
        for w in words:
            if c.stringWidth(curr_line + " " + w, "Helvetica", 10) < (width - 108):
                curr_line = (curr_line + " " + w).strip()
            else:
                c.drawString(54, y, curr_line)
                y -= line_height
                curr_line = w
                if y < margin_bottom:
                    c.showPage()
                    c.setFont("Helvetica", 10)
                    y = height - 54
        
        if curr_line:
            c.drawString(54, y, curr_line)
            y -= line_height

        if y < margin_bottom:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = height - 54

    c.save()
    bio.seek(0)
    return bio
